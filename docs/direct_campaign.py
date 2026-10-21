"""Кампания Яндекс Директа для лендинга /port → docs/direct-campaign.xlsx. Запуск: python docs/direct_campaign.py"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

F = "Arial"
wb = Workbook()
H = Font(name=F, bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="0B5C8F")
B = Font(name=F); BB = Font(name=F, bold=True); INP = Font(name=F, color="0000FF")
YEL = PatternFill("solid", fgColor="FFF2B3"); RED = PatternFill("solid", fgColor="F8D0CC")

def head(ws, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        x = ws.cell(1, i, c); x.font = H; x.fill = HF; x.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

# --- 1. Настройки
ws = wb.active; ws.title = "Настройки"
ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 70; ws.column_dimensions["C"].width = 60
rows = [
 ("Параметр", "Значение", "Пояснение"),
 ("Адрес лендинга", "https://renal-auto.ru/port", "Жёлтые ячейки можно менять. Ссылки на листе «Объявления» собираются из этого адреса формулой. Запускать рекламу — после переезда на renal-auto.ru (docs/MOVE_SERVER.md)."),
 ("UTM-метки", "utm_source=yandex&utm_medium=cpc&utm_campaign={campaign_id}&utm_content={ad_id}&utm_term={keyword}", "{…} — динамические параметры Директа, подставятся сами."),
 ("Тип кампании", "Единая перфоманс-кампания", "Для старта — только поиск: в «Местах показа» отключите сети, пока не наберётся статистика."),
 ("Регион показа", "Россия", "Допущение: возим по всей России. Если хотите начать с Дальнего Востока и Сибири — сузьте здесь."),
 ("Стратегия на первые 2 недели", "Максимум кликов, недельный бюджет", "После 10+ заявок в неделю — переключить на «Максимум конверсий» по цели «lead»."),
 ("Недельный бюджет теста, ₽", 15000, "Допущение для теста, не рекомендация по итоговому бюджету — поменяйте под себя."),
 ("Счётчик Метрики", 113699503, "Уже стоит на лендинге."),
 ("Цель для оптимизации", "lead — отправленная заявка", "Другие цели лендинга: car_click, car_trip, form_open, max, reviews_2gis — для анализа, не для оптимизации."),
 ("Расписание", "Ежедневно 8:00–23:00", "По времени региона показа."),
 ("Минус-слова", "лист «Минус-слова»", "Добавить на уровень кампании целиком."),
 ("Быстрые ссылки и уточнения", "лист «Быстрые ссылки»", "Один набор на все группы."),
 ("Проверка длины", "столбцы «Длина» на листе «Объявления»", "Красная заливка — текст длиннее лимита Директа (заголовок 1 — 56, заголовок 2 — 30, текст — 81, отображаемая ссылка — 20)."),
]
for r, row in enumerate(rows, 1):
    for c, v in enumerate(row, 1):
        x = ws.cell(r, c, v); x.font = BB if r == 1 or c == 1 else B; x.alignment = Alignment(wrap_text=True, vertical="top")
        if r == 1: x.font = H; x.fill = HF
for r in (2, 3, 5, 7): ws.cell(r, 2).fill = YEL; ws.cell(r, 2).font = INP

# --- группы: (кампания, группа, модель для ?model=, название в объявлении, отображаемая ссылка, русские написания)
KR, CN = "Корея — модели", "Китай — модели"
groups = [
 (KR, "Kia K5", "K5", "Kia K5", "Kia-K5", ["киа к5", "kia k5", "киа оптима к5"]),
 (KR, "Kia Sportage", "Sportage", "Kia Sportage", "Kia-Sportage", ["киа спортейдж", "kia sportage"]),
 (KR, "Kia Sorento", "Sorento", "Kia Sorento", "Kia-Sorento", ["киа соренто", "kia sorento"]),
 (KR, "Kia Seltos", "Seltos", "Kia Seltos", "Kia-Seltos", ["киа селтос", "kia seltos"]),
 (KR, "Kia Carnival", "Carnival", "Kia Carnival", "Kia-Carnival", ["киа карнивал", "kia carnival"]),
 (KR, "Kia Ray", "Ray", "Kia Ray", "Kia-Ray", ["киа рей", "kia ray"]),
 (KR, "Hyundai Tucson", "Tucson", "Hyundai Tucson", "Hyundai-Tucson", ["хендай туссан", "хендай туксон", "hyundai tucson"]),
 (KR, "Hyundai Santa Fe", "Santa Fe", "Hyundai Santa Fe", "Hyundai-Santa-Fe", ["хендай санта фе", "hyundai santa fe"]),
 (KR, "Hyundai Sonata", "Sonata", "Hyundai Sonata", "Hyundai-Sonata", ["хендай соната", "hyundai sonata"]),
 (KR, "Hyundai Avante", "Avante", "Hyundai Avante", "Hyundai-Avante", ["хендай аванте", "hyundai avante", "хендай элантра из кореи"]),
 (KR, "Genesis G80", "G80", "Genesis G80", "Genesis-G80", ["генезис g80", "genesis g80"]),
 (CN, "Geely Monjaro", "Monjaro", "Geely Monjaro", "Geely-Monjaro", ["джили монжаро", "geely monjaro"]),
 (CN, "Geely Coolray", "Coolray", "Geely Coolray", "Geely-Coolray", ["джили кулрей", "geely coolray"]),
 (CN, "Haval Jolion", "Jolion", "Haval Jolion", "Haval-Jolion", ["хавал джолион", "haval jolion"]),
 (CN, "Haval Dargo", "Dargo", "Haval Dargo", "Haval-Dargo", ["хавал дарго", "haval dargo"]),
 (CN, "Chery Tiggo 7 Pro", "Tiggo 7", "Chery Tiggo 7 Pro", "Chery-Tiggo-7", ["чери тигго 7", "chery tiggo 7"]),
 (CN, "Chery Tiggo 8", "Tiggo 8", "Chery Tiggo 8", "Chery-Tiggo-8", ["чери тигго 8", "chery tiggo 8"]),
 (CN, "Changan CS75 Plus", "CS75", "Changan CS75 Plus", "Changan-CS75", ["чанган cs75", "changan cs75"]),
 (CN, "Changan UNI-K", "UNI-K", "Changan UNI-K", "Changan-UNI-K", ["чанган юни к", "changan uni k"]),
]
general = [
 ("Общие", "Авто из Кореи", "", "Авто из Кореи", "Авто-из-Кореи"),
 ("Общие", "Авто из Китая", "", "Авто из Китая", "Авто-из-Китая"),
 ("Общие", "Подбор авто из Кореи и Китая", "", "Авто из Кореи и Китая", "Подбор-авто"),
]

# --- 2. Объявления: два варианта на группу
ws = wb.create_sheet("Объявления")
cols = ["Кампания", "Группа", "Вариант", "Заголовок 1", "Длина", "Заголовок 2", "Длина", "Текст", "Длина", "Отображаемая ссылка", "Длина", "?model=", "Ссылка (формула)"]
head(ws, cols, [16, 22, 8, 48, 7, 26, 7, 62, 7, 20, 7, 12, 70])
def ad_rows(camp, grp, model, name, disp, country):
    from_ = {"KR": "из Кореи", "CN": "из Китая", "": "из Кореи и Китая"}[country]
    t1a = f"{name} {from_} — подберём оптимальный вариант" if country else f"{name} — подберём оптимальный вариант"
    if len(t1a) > 56: t1a = f"{name} {from_} — подберём вариант"
    if len(t1a) > 56: t1a = f"{name} {from_}"
    t1b = f"{name} {from_}: проверим историю до покупки" if country else f"{name}: проверим историю до покупки"
    if len(t1b) > 56: t1b = f"{name} {from_}: проверка до покупки"
    if len(t1b) > 56: t1b = f"{name}: проверка до покупки"
    return [
        (camp, grp, "А", t1a, "Комиссия до 50 000 ₽", "Сравним объявления, проверим историю и проходимость в РФ. Цены онлайн в рублях.", disp, model),
        (camp, grp, "Б", t1b, "Агентский договор", "Реальные объявления с ценой в рублях. Посмотрите, как машина поедет к вам.", disp, model),
    ]
allrows = []
for camp, grp, model, name, disp, _ in groups:
    allrows += ad_rows(camp, grp, model, name, disp, "KR" if camp == KR else "CN")
for camp, grp, model, name, disp in general:
    allrows += ad_rows(camp, grp, model, name, disp, "")  # «из Кореи» уже в названии группы
for i, (camp, grp, v, t1, t2, txt, disp, model) in enumerate(allrows, 2):
    vals = [camp, grp, v, t1, f"=LEN(D{i})", t2, f"=LEN(F{i})", txt, f"=LEN(H{i})", disp, f"=LEN(J{i})", model,
            f'=Настройки!$B$2&IF(L{i}="","?","?model="&SUBSTITUTE(L{i}," ","+")&"&")&Настройки!$B$3']
    for c, val in enumerate(vals, 1):
        x = ws.cell(i, c, val); x.font = B; x.alignment = Alignment(wrap_text=c in (4, 8, 13), vertical="top")
n = len(allrows) + 1
for col, lim in (("E", 56), ("G", 30), ("I", 81), ("K", 20)):
    ws.conditional_formatting.add(f"{col}2:{col}{n}", FormulaRule(formula=[f"{col}2>{lim}"], fill=RED))

# --- 3. Фразы
ws = wb.create_sheet("Фразы")
head(ws, ["Кампания", "Группа", "Ключевая фраза"], [16, 24, 50])
r = 2
for camp, grp, model, name, disp, ru in groups:
    land = "кореи" if camp == KR else "китая"
    for base in ru:
        for tpl in ("{b} из {l}", "купить {b} из {l}", "{b} из {l} цена", "заказать {b} из {l}", "пригнать {b} из {l}"):
            if " из " in base: phrase = base if tpl == "{b} из {l}" else None
            else: phrase = tpl.format(b=base, l=land)
            if phrase:
                for c, v in enumerate((camp, grp, phrase), 1): ws.cell(r, c, v).font = B
                r += 1
gen = {
 "Авто из Кореи": ["авто из кореи", "купить авто из кореи", "заказать авто из кореи", "пригнать авто из кореи", "машина из кореи", "авто из кореи под ключ", "доставка авто из кореи", "автомобили из южной кореи"],
 "Авто из Китая": ["авто из китая", "купить авто из китая", "заказать авто из китая", "пригнать авто из китая", "машина из китая", "авто из китая под ключ", "доставка авто из китая"],
 "Подбор авто из Кореи и Китая": ["подбор авто из кореи", "подбор авто из китая", "авто из кореи и китая", "растаможка авто из кореи", "расчет авто из кореи", "калькулятор авто из кореи", "калькулятор растаможки авто из китая"],
}
for grp, ps in gen.items():
    for p in ps:
        for c, v in enumerate(("Общие", grp, p), 1): ws.cell(r, c, v).font = B
        r += 1
ws.cell(r + 1, 1, "Фразы с «под ключ» оставлены: так ищут. В текстах объявлений «под ключ» не используем — цена ориентировочная.").font = Font(name=F, italic=True)

# --- 4. Минус-слова
ws = wb.create_sheet("Минус-слова")
head(ws, ["Минус-слово", "Зачем"], [26, 60])
neg = [("запчасти", "не продаём запчасти"), ("разборка", ""), ("разбор", ""), ("ремонт", ""), ("сервис", ""), ("шины", ""), ("диски", ""), ("тюнинг", ""),
       ("аренда", "не сдаём машины"), ("прокат", ""), ("такси", ""), ("лизинг", "не оформляем лизинг — уберите, если появится"),
       ("вакансии", "не ищем сотрудников"), ("работа", ""), ("водитель", ""),
       ("игрушка", "модели и игры"), ("масштабная", ""), ("гта", ""), ("игра", ""), ("скачать", ""), ("мод", ""), ("раскраска", ""),
       ("обзор", "ищут информацию, а не покупку"), ("тест-драйв", ""), ("характеристики", ""), ("расход", ""), ("габариты", ""),
       ("японии", "не возим из Японии"), ("сша", "не возим из США"), ("европы", ""), ("германии", ""), ("оаэ", ""), ("грузии", ""),
       ("авито", "ищут на площадках, не у нас"), ("дром", ""), ("автору", "")]
for i, (w, why) in enumerate(neg, 2):
    ws.cell(i, 1, w).font = B; ws.cell(i, 2, why).font = B

# --- 5. Быстрые ссылки и уточнения
ws = wb.create_sheet("Быстрые ссылки")
head(ws, ["Тип", "Заголовок", "Длина", "Описание", "Длина", "Якорь на странице", "Ссылка (формула)"], [14, 28, 7, 58, 7, 18, 60])
sl = [("Быстрая ссылка", "Объявления онлайн", "Свежие машины из Кореи и Китая с ценой в рублях", "#terminal"),
      ("Быстрая ссылка", "Как машина поедет к вам", "Покажем путь от площадки до вашего города", "#route"),
      ("Быстрая ссылка", "Отзывы и выдачи", "Видео с выдач и отзывы клиентов в 2ГИС", "#vydachi"),
      ("Быстрая ссылка", "Проверьте нас", "Юрлицо, офис во Владивостоке, договор", "#vydachi")]
for i, (t, title, desc, anchor) in enumerate(sl, 2):
    for c, v in enumerate((t, title, f"=LEN(B{i})", desc, f"=LEN(D{i})", anchor, f"=Настройки!$B$2&\"?\"&Настройки!$B$3&F{i}"), 1):
        ws.cell(i, c, v).font = B
r = len(sl) + 2
for u in ("Агентский договор", "Комиссия до 50 000 ₽", "Проверка истории и VIN", "Видео с площадки"):
    for c, v in enumerate(("Уточнение", u, f"=LEN(B{r})"), 1): ws.cell(r, c, v).font = B
    r += 1
ws.conditional_formatting.add(f"C2:C{len(sl)+1}", FormulaRule(formula=["C2>30"], fill=RED))
ws.conditional_formatting.add(f"E2:E{len(sl)+1}", FormulaRule(formula=["E2>60"], fill=RED))
ws.conditional_formatting.add(f"C{len(sl)+2}:C{r-1}", FormulaRule(formula=[f"C{len(sl)+2}>25"], fill=RED))

out = "/home/user/RenalAuto/docs/direct-campaign.xlsx"
wb.save(out); print(out, len(allrows), "объявлений")
